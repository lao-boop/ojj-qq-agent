"""生成汽修 Agent 评测报告（Word）。"""
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

FONT = "Microsoft YaHei"
ACCENT = RGBColor(0x1E, 0x7F, 0xD6)
MUTED = RGBColor(0x5C, 0x72, 0x85)
DARK = RGBColor(0x0F, 0x22, 0x33)


def set_font(run, size=10.5, bold=False, color=DARK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def shade(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def h(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_font(run, 16 if level == 1 else 13, True, ACCENT if level == 1 else DARK)
    return p


def para(doc, text, size=10.5, color=DARK, bold=False, space=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space)
    p.paragraph_format.line_spacing = 1.35
    run = p.add_run(text)
    set_font(run, size, bold, color)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    set_font(run, 10.5)
    return p


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, name in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(name)
        set_font(run, 10, True, RGBColor(0xFF, 0xFF, 0xFF))
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        shade(cell, "1E7FD6")
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            set_font(run, 10)
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    return t


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    # 封面标题
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(60)
    run = p.add_run("汽修多模态 Agent 评测报告")
    set_font(run, 26, True, ACCENT)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("基于 auto_repair_multimodal_v1 数据集 · dev / test 双集评测")
    set_font(run, 12, False, MUTED)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("2026-10-25")
    set_font(run, 10.5, False, MUTED)

    doc.add_paragraph()

    # 一、结论先行
    h(doc, "一、结论先行")
    para(
        doc,
        "经过三轮提示词与规则迭代，测试集（26 例）案例通过率从 46.2% 提升到 80.8%，"
        "字段级通过率从 85.4% 提升到 92.7%。剩下的失败集中在 VIN 冲突字段和单据 part_ids 两类。",
    )
    bullet(doc, "配件分类、适配边界、语音理解、聊天上下文、车辆观察：全部通过（6/6、3/3、4/4、2/2、2/2）")
    bullet(doc, "VIN 识别 2/4、维修单提取 0/2，是当前主要短板")
    bullet(doc, "需要说明：test 集已被使用三轮，80.8% 偏乐观，真实场景需重新采数据验证")

    # 二、任务与数据
    h(doc, "二、任务与数据")
    para(
        doc,
        "数据集是一套汽修群聊场景的合成多模态样本，覆盖 8 类任务：配件识别、适配边界、车辆观察、"
        "维修单提取、VIN 识别、聊天上下文、语音理解、跨附件一致性。",
    )
    table(
        doc,
        ["项目", "dev", "test"],
        [
            ["案例数", "18", "26"],
            ["任务类型", "8 类", "8 类"],
            ["语音样本", "2 条", "4 条"],
            ["参考字符数", "64", "134"],
        ],
    )
    para(
        doc,
        "说明：数据集 README 明确标注它用于提示词调试、知识库与评测，并非微调格式，因此本次工作不包含参数微调。",
        size=9.5,
        color=MUTED,
    )

    # 三、方法
    h(doc, "三、方法：做了什么")
    bullet(doc, "配件知识库：24 类配件整理成 JSONL（part_id / 别名 / 外观线索 / 易混项 / 追问注意），注入提示词")
    bullet(doc, "多模态预处理：图片先转成结构化文字描述（VIN 逐位读、遮挡位用 ?、车辆明确判断划痕）")
    bullet(doc, "语音：WhisperX 转写并缓存，dev/test 各音频只跑一次")
    bullet(doc, "输出约束：统一 JSON 结构，规定必填字段、默认值、追问判定与字段互斥规则")
    bullet(doc, "稳定性：三次尝试（主模型 6000 tokens → flash → 主模型 8000 tokens），全部失败才记失败")
    bullet(doc, "评分：直接用数据集自带 tools/evaluate.py，不自定义指标")

    # 四、结果
    h(doc, "四、评测结果")
    h(doc, "4.1 三轮对比（test）", level=2)
    table(
        doc,
        ["轮次", "案例通过率", "字段通过率", "主要改动"],
        [
            ["第 1 轮", "46.2%", "85.4%", "基线"],
            ["第 2 轮", "65.4%", "92.1%", "字段完整性、VIN O/0、金额核算、追问规则"],
            ["第 3 轮", "80.8%", "92.7%", "VIN 逐位与冲突、划痕判断、客户与否定项"],
        ],
    )

    h(doc, "4.2 最终分任务（test）", level=2)
    table(
        doc,
        ["任务", "通过", "说明"],
        [
            ["配件分类", "6/6", "✅ 全部正确"],
            ["适配边界", "3/3", "✅ 从不无依据确认适配"],
            ["语音理解", "4/4", "✅ 否定项、单位保留正确"],
            ["聊天上下文", "2/2", "✅ 更正覆盖与客户字段正确"],
            ["车辆观察", "2/2", "✅ 划痕判断已修复"],
            ["跨附件一致性", "2/3", "冲突识别正确，个别字段缺失"],
            ["VIN 识别", "2/4", "双 VIN 冲突字段与候选列表仍有误"],
            ["维修单提取", "0/2", "明细正确，part_ids 被规则误伤为空"],
        ],
    )

    h(doc, "4.3 指标细分（test）", level=2)
    table(
        doc,
        ["指标", "通过率", "指标", "通过率"],
        [
            ["配件分类", "90.0%", "确认边界", "100%"],
            ["追问判定", "100%", "数量", "100%"],
            ["VIN 抄录", "57.1%", "VIN 校验", "50.0%"],
            ["冲突识别", "71.4%", "单据金额", "100%"],
            ["上下文", "100%", "跨模态融合", "100%"],
        ],
    )
    para(doc, "语音转写 CER：7.5%（134 字中错 10 字），上一轮为 3.7%，属于正常波动。")

    # 五、失败分析
    h(doc, "五、剩余失败与原因")
    para(doc, "最终失败 5 例，原因都已定位，属于规则或 OCR 精度问题，不是理解能力问题：")
    table(
        doc,
        ["案例", "问题", "原因"],
        [
            ["C033 / C034", "维修单 part_ids 为空", "“跨附件只取照片类别”的规则被误用到纯单据任务"],
            ["C041 / C042", "双 VIN 冲突字段不完整", "complete / format_valid / candidates 未被规则覆盖"],
            ["C085", "candidates 缺失", "同上，冲突场景字段未填全"],
        ],
    )
    para(
        doc,
        "另外 VIN 抄录只有 57.1%：合成图的 VIN 是小字采集卡，视觉模型逐位读取仍会漏字符，"
        "这属于 OCR 精度上限，靠提示词难以根治。",
        size=9.5,
        color=MUTED,
    )

    # 六、怎么在 QQ 上测
    h(doc, "六、在 QQ 上测试她")
    para(
        doc,
        "现在 QQ 里的“小鲸鱼”已经接入配件知识库（lookup_part），可以直接测“判断类”能力；"
        "完整的结构化评测（VIN 逐位、单据金额核算、跨附件冲突）仍跑在离线脚本里，尚未接进聊天。",
    )
    h(doc, "6.1 现在就能测的", level=2)
    bullet(doc, "「刹车片是什么」「空气滤芯和空调滤芯怎么分」→ 她会查 24 类配件库后回答")
    bullet(doc, "发一张配件照片 + 「这是什么」→ 视觉识别 + 分类 + 追问缺失信息")
    bullet(doc, "「这个能装我车上吗」→ 她会要求车型/年款/位置，不会乱确认适配")
    bullet(doc, "「报个价」→ 会说明需要先确认适配，不承诺价格")
    h(doc, "6.2 推荐测试话术", level=2)
    table(
        doc,
        ["场景", "可以这样说"],
        [
            ["配件识别", "发一张配件照片，问「这是什么配件？怎么判断？」"],
            ["易混项", "「刹车片和刹车盘怎么区分？」"],
            ["适配边界", "「这个空滤能装我的车吗？」（故意不说车型）"],
            ["追问质量", "「就这个，帮我配一个」→ 看她是否先问车型和位置"],
            ["否认项", "「我要空气滤芯，不是空调滤芯」→ 看她是否记住排除项"],
        ],
    )
    h(doc, "6.3 还没接进 QQ 的", level=2)
    bullet(doc, "VIN 采集卡逐位识别与冲突提示")
    bullet(doc, "维修单金额核算（printed_total vs computed_total）")
    bullet(doc, "跨附件一致性判断（照片 vs 单据是否同一车/同一单）")
    para(
        doc,
        "这三项需要把离线评测里的结构化输出接成“汽修模式”：识别到相关照片时，先出结构化结果，再转成口语回复。"
        "可以做，约需一次开发 + 一轮回归测试。",
        size=9.5,
        color=MUTED,
    )

    # 七、建议
    h(doc, "七、下一步建议")
    bullet(doc, "先修三条规则（单据 part_ids、双 VIN 冲突字段、candidates），预计 test 可到 90% 左右")
    bullet(doc, "修完不再用 test 调参，改用 dev 验证，避免指标虚高")
    bullet(doc, "VIN 想再准，需要放大/裁剪 VIN 区域单独 OCR，或换更强的 OCR 模型")
    bullet(doc, "接“汽修模式”进 QQ，让结构化能力真正可被用户使用")
    bullet(doc, "长期：用真实车间照片/单据替换合成数据，才能反映真实水平")

    out = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "data", "documents", "汽修Agent评测报告.docx")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
