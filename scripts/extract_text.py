import sys
import json
import os

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

ASCII = True  # JSON 里中文用 \\uXXXX 转义，避免 Windows 控制台编码问题


def emit(obj):
    print(json.dumps(obj, ensure_ascii=ASCII))


def main():
    if len(sys.argv) < 2:
        emit({"ok": False, "error": "missing path"})
        return
    p = sys.argv[1]
    ext = os.path.splitext(p)[1].lower()
    text = ""
    try:
        if ext == ".docx":
            import docx
            d = docx.Document(p)
            parts = [x.text for x in d.paragraphs if x.text and x.text.strip()]
            for t in d.tables:
                for row in t.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))
            text = "\n".join(parts)
        elif ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(p)
            chunks = []
            for i, slide in enumerate(prs.slides, 1):
                lines = ["第%d页:" % i]
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False):
                        s = shape.text_frame.text.strip()
                        if s:
                            lines.append(s)
                chunks.append("\n".join(lines))
            text = "\n".join(chunks)
        elif ext in (".xlsx", ".xlsm"):
            import openpyxl
            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
            chunks = []
            for ws in wb.worksheets:
                chunks.append("[工作表 %s]" % ws.title)
                for row in ws.iter_rows(values_only=True):
                    vals = [str(v) for v in row if v is not None]
                    if vals:
                        chunks.append(" | ".join(vals))
            text = "\n".join(chunks)
        elif ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(p)
            chunks = []
            for i, page in enumerate(reader.pages, 1):
                page_text = page.extract_text() or ""
                if page_text.strip():
                    chunks.append("第%d页: %s" % (i, page_text))
            text = "\n".join(chunks)
        elif ext in (".txt", ".md", ".csv", ".json", ".log", ".py", ".js", ".ts", ".html", ".xml", ".yml", ".yaml"):
            raw = open(p, "rb").read()
            for enc in ("utf-8", "gbk", "utf-16"):
                try:
                    text = raw.decode(enc)
                    break
                except Exception:
                    continue
        elif ext in (".doc", ".xls", ".ppt"):
            emit({"ok": False, "error": "旧版 Office 格式（.doc/.xls/.ppt）不支持，请另存为 docx/xlsx/pptx"})
            return
        else:
            emit({"ok": False, "error": "不支持的格式 %s" % ext})
            return

        text = text.strip()
        emit({"ok": True, "chars": len(text), "text": text[:20000]})
    except Exception as e:
        emit({"ok": False, "error": str(e)})


if __name__ == "__main__":
    main()
