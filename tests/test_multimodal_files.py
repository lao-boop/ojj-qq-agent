import contextlib
import importlib.util
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import math
import struct
import wave
from unittest.mock import patch

ROOT = Path(__file__).resolve().parent.parent


class FilesTest(unittest.TestCase):
    def run_script(self, name, *args):
        result = subprocess.run([sys.executable, str(ROOT / 'scripts' / name), *map(str, args)], check=True, capture_output=True, text=True)
        return json.loads(result.stdout.strip().splitlines()[-1])

    def test_real_silk_decode_standard_and_tencent_headers(self):
        import pilk
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            pcm = root / 'voice.pcm'
            pcm.write_bytes(b''.join(struct.pack('<h', int(8000 * math.sin(2 * math.pi * 440 * i / 24000))) for i in range(24000)))
            for tencent in (False, True):
                silk, output = root / f'{tencent}.silk', root / f'{tencent}.wav'
                pilk.encode(str(pcm), str(silk), pcm_rate=24000, tencent=tencent)
                subprocess.run([sys.executable, str(ROOT / 'scripts' / 'silk_to_wav.py'), str(silk), str(output), '24000'], check=True, capture_output=True)
                with wave.open(str(output)) as decoded:
                    self.assertEqual(decoded.getframerate(), 24000)
                    self.assertEqual(decoded.getnchannels(), 1)
                    self.assertGreater(decoded.getnframes(), 22000)

    def test_office_text_and_scanned_pdf(self):
        from docx import Document
        from openpyxl import Workbook
        from pptx import Presentation
        from pypdf import PdfWriter
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            doc = Document(); doc.add_paragraph('项目交付日期：2026-11-01'); doc.save(root / 'task.docx')
            book = Workbook(); book.active.append(['项目', '进度']); book.active.append(['任务甲', '80%']); book.save(root / 'task.xlsx')
            slides = Presentation(); slide = slides.slides.add_slide(slides.slide_layouts[1]); slide.shapes.title.text = '项目验收'; slides.save(root / 'task.pptx')
            for name, expected in [('task.docx', '2026-11-01'), ('task.xlsx', '80%'), ('task.pptx', '项目验收')]:
                result = self.run_script('extract_text.py', root / name)
                self.assertTrue(result['ok'], result)
                self.assertIn(expected, result['text'])
            writer = PdfWriter()
            for _ in range(4): writer.add_blank_page(width=300, height=400)
            writer.write(root / 'scan.pdf')
            text = self.run_script('extract_text.py', root / 'scan.pdf')
            self.assertTrue(text['ok']); self.assertEqual(text['text'], '')
            media = self.run_script('extract_media.py', root / 'scan.pdf', root / 'pages')
            self.assertTrue(media['ok'], media)
            self.assertEqual(len(media['images']), 3)
            self.assertTrue(all(Path(p).stat().st_size > 300 for p in media['images']))

    def test_config_wizard_preserves_accounts_and_writes_private_backup(self):
        spec = importlib.util.spec_from_file_location('configure', ROOT / 'scripts' / 'configure_multimodal.py')
        module = importlib.util.module_from_spec(spec); spec.loader.exec_module(module)
        with tempfile.TemporaryDirectory() as tmp:
            home = Path(tmp)
            shared = home / 'qq-agent-dev' / 'napcat' / 'data' / 'qq'; shared.mkdir(parents=True)
            cfg = {'onebot': {'accessToken': 'keep-qq'}, 'llm': {'model': 'keep-model', 'apiKey': 'keep-text-key'}, 'access': {'ownerQQ': '10001'}, 'image': {'enabled': False}, 'asr': {'local': {'model': 'medium'}}}
            target = home / 'config.json'; target.write_text(json.dumps(cfg))
            outputs = io.StringIO()
            with patch.dict(os.environ, {'QQBOT_CONFIG': str(target)}), patch.object(module.Path, 'home', return_value=home), patch('builtins.input', return_value=''), patch.object(module.getpass, 'getpass', return_value='fake-minimax-key'), contextlib.redirect_stdout(outputs):
                module.main()
            saved = json.loads(target.read_text())
            self.assertEqual(saved['onebot'], cfg['onebot'])
            self.assertEqual(saved['llm'], cfg['llm'])
            self.assertEqual(saved['access'], cfg['access'])
            self.assertEqual(saved['tts']['provider'], 'minimax')
            self.assertEqual(saved['tts']['minimax']['apiKey'], 'fake-minimax-key')
            self.assertFalse(saved['vision']['enabled'])
            self.assertTrue(saved['chat']['agents']['enabled'])
            self.assertEqual(saved['asr']['local']['model'], 'medium')
            self.assertEqual(target.stat().st_mode & 0o777, 0o600)
            backup = next((home / '.qq-agent-backups').iterdir())
            self.assertEqual(json.loads(backup.read_text()), cfg)
            self.assertEqual(backup.stat().st_mode & 0o777, 0o600)
            self.assertNotIn('fake-minimax-key', outputs.getvalue())


if __name__ == '__main__':
    unittest.main()
